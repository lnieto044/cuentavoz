# -*- coding: utf-8 -*-
"""Arma el guion de la presentacion en Word, y lo mete en las notas del
orador del PowerPoint.

Dos salidas de una sola fuente, para que no puedan desincronizarse:

  · docs/Guion_Presentacion.docx  - para imprimir y ensayar
  · las notas del orador del .pptx - lo que ven en la vista de presentador
    mientras exponen, sin tener que mirar una hoja aparte

El contenido vive aqui, en GUION, y no en el .docx ni en el .pptx: editar
un binario a mano es como se desincronizan las cosas.

    python docs/armar_guion.py
"""
import io
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm
from pptx import Presentation

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
DOCX = os.path.join(AQUI, "Guion_Presentacion.docx")
PPTX = os.path.join(AQUI, "CuentaVoz_Colsubsidio_V2.pptx")

# Paleta oficial de Colsubsidio, la misma de la app y los manuales
AZUL = RGBColor(0x00, 0x67, 0xB1)
NAVY = RGBColor(0x10, 0x29, 0x4C)
GRAFITO = RGBColor(0x57, 0x57, 0x56)
AMARILLO_TX = RGBColor(0x8A, 0x6D, 0x00)

# (nº, titulo, quien, segundos, [lineas del guion], [notas de puesta en escena])
GUION = [
    (1, "Portada", "LUIS", 20,
     ["Buenos días. Somos StockXperts. Yo soy Luis, ella es Diana.",
      "Traemos CuentaVoz: contar el inventario hablando, en vez de escribiendo.",
      "En diez minutos se lo mostramos funcionando con los datos reales de ustedes."],
     ["No lea la diapositiva: ya se ve."]),

    (2, "El conteo manual cuesta caro", "LUIS", 40,
     ["Hoy el inventario de una bodega se toma en papel. Alguien camina con una "
      "planilla, anota, y después otra persona lo digita en el sistema.",
      "Eso tiene tres costos: TIEMPO —se hace dos veces—, ERRORES —de letra, de "
      "digitación— y DEMORA: la diferencia se descubre semanas después, cuando ya "
      "nadie se acuerda de qué pasó."],
     ["Pausa corta antes de pasar."]),

    (3, "La solución, en una frase", "LUIS", 30,
     ["Voz, inteligencia artificial y validaciones automáticas.",
      "La persona habla, el sistema entiende, valida contra el inventario y guarda.",
      "Una sola vez, en el sitio, y validado en el momento."],
     []),

    (4, "Cómo piensa el agente", "DIANA", 40,
     ["Cada vez que alguien habla pasan cinco pasos: escucha, interpreta, resuelve "
      "el artículo contra el catálogo oficial, valida y confirma.",
      "Y algo importante: el modelo interpreta, pero el que decide es el backend. "
      "La inteligencia artificial no escribe en la base de datos — propone, y unas "
      "reglas en código deciden si eso entra o no."],
     ["Ese matiz tranquiliza a quien pregunta por confiabilidad de la IA."]),

    (5, "Tres momentos, una sola plataforma", "DIANA", 40,
     ["No es solo contar. Son los tres momentos del ciclo: el PEDIDO al almacén "
      "calculado por receta, el CONTEO de la bodega, y la LEGALIZACIÓN del "
      "servicio al final del turno.",
      "Todo sin salir de la misma aplicación, con la misma voz."],
     ["Al terminar esta lámina deberían ir en 2:50."]),

    (6, "DEMO EN VIVO", "LUIS cuenta · DIANA cierra", 120,
     ["LUIS (1:10) — como auxiliar:",
      "1. Entro como auxiliar. Solo veo MIS bodegas asignadas, ninguna más.",
      "2. Abre una bodega. «Ya está abierta, con sus 133 referencias.»",
      "3. Toca el micrófono y dicta despacio: «papa criolla veinte kilos».",
      "4. Me repite lo que entendió y me pide confirmar. → confirma.",
      "5. Si sale diferencia: «Aquí me avisa que el sistema esperaba otra cantidad. "
      "ESTO es el hallazgo: se ve ahora, no en tres semanas.»",
      "6. Cuando termino, firmo mi conteo. → «Terminar y firmar mi conteo».",
      "",
      "DIANA (0:50) — como administradora:",
      "7. Yo recibo esa bodega firmada y hago el RECUENTO CIEGO: cuento sin ver sus "
      "números.",
      "8. Muestra la comparación: «aquí se revelan las tres columnas: el sistema, lo "
      "que contó Luis, lo que conté yo».",
      "9. Y se cierra con DOBLE FIRMA. Ninguno de los dos puede cerrar solo."],
     ["NO use las capturas: abra la tableta.",
      "Si el micrófono falla, toque Teclado y escriba la misma frase. Es el "
      "respaldo real de la aplicación, no una excusa: dígalo así, «funciona en una "
      "bodega con ruido».",
      "CORTE: si a los 5:00 no han salido de la demo, córtenla y sigan."]),

    (7, "Nada se guarda sin validar", "LUIS", 40,
     ["Lo que acaban de ver validó cuatro cosas sin que yo hiciera nada: que el "
      "artículo existe en el catálogo, que la unidad es la correcta, que la "
      "cantidad tiene sentido, y que yo tenía permiso sobre esa bodega.",
      "Si algo no cuadra, el conteo NO se detiene: queda marcado y sigue."],
     []),

    (8, "Arquitectura", "DIANA", 40,
     ["Está en AWS: el frontend en S3 con CloudFront, el backend en EC2 con Docker, "
      "la base en RDS. La misma nube que ya usa Colsubsidio, así que no hay nada "
      "que migrar.",
      "Y cada cambio pasa por 149 pruebas automáticas antes de desplegarse. Si una "
      "falla, no llega a producción."],
     ["Es la lámina donde más preguntan. No se extienda: si preguntan, hay tiempo "
      "al final."]),

    (9, "Datos reales", "DIANA", 30,
     ["Nada de esto es una maqueta: son las 54 bodegas de ustedes, 1.041 artículos "
      "del catálogo oficial y 1.405 registros de stock.",
      "Al cargarlo, el sistema encontró 79 saldos negativos — inventario que el "
      "sistema dice tener en menos que cero."],
     []),

    (10, "Seguridad", "LUIS", 30,
     ["La identidad la maneja AWS Cognito, con verificación en dos pasos opcional. "
      "La clave nunca pasa por nuestro servidor.",
      "Y nada se borra: una corrección crea un registro nuevo que apunta al "
      "original. La trazabilidad es completa."],
     ["Al terminar esta lámina deberían ir en 7:10."]),

    (11, "Qué significa para Colsubsidio", "LUIS", 30,
     ["Decisiones con datos reales, no con memoria ni planillas sueltas.",
      "El auxiliar recupera tiempo, el administrador ve el estado en vivo."],
     ["Una sola idea. No enumere las tres viñetas: escoja la que más le importe a "
      "quien tiene enfrente."]),

    (12, "Costo", "DIANA", 30,
     ["No hay licencias por usuario. Sumar una bodega o veinte personas más no "
      "cuesta más. Se paga la infraestructura y el consumo de la IA por uso."],
     []),

    (13, "Adopción", "DIANA", 30,
     ["Hoy: prototipo real desplegado y funcionando.",
      "Próximo paso: cargar el archivo oficial de recetas y un piloto guiado en un "
      "grupo de bodegas.",
      "Visión: toda la operación de hotelería, integrada con My Inventory."],
     []),

    (14, "Qué necesitamos hoy", "LUIS", 40,
     ["Para pasar del prototipo al piloto necesitamos cuatro cosas concretas: un "
      "SPONSOR dentro de Operaciones, un GRUPO DE BODEGAS para el piloto, el "
      "ARCHIVO OFICIAL DE RECETAS, y ACCESO A MY INVENTORY para cerrar el ciclo."],
     ["ESTA LÁMINA NO SE SACRIFICA. Si va corto, recorte la 11 o la 12, nunca esta.",
      "Mire a la persona que puede decir que sí."]),

    (15, "Equipo", "DIANA", 15,
     ["Somos el equipo StockXperts. Esto lo construimos nosotros, de cero."],
     []),

    (16, "Gracias", "LUIS", 15,
     ["Ahí está la dirección para que lo prueben ustedes mismos, con estas cuentas. "
      "Quedamos atentos a sus preguntas."],
     ["Deje esta lámina puesta durante las preguntas: tiene la URL, las cuentas y "
      "el contacto."]),
]

PREGUNTAS = [
    ("¿Qué pasa si no hay señal en la bodega?", "LUIS",
     "El conteo no se detiene. Se guarda en la tableta y se sincroniza solo al "
     "volver la conexión. Es una PWA real."),
    ("¿Y si la IA entiende mal?", "DIANA",
     "Confirma antes de guardar, siempre. Y si el nombre se parece a varios "
     "artículos, muestra las opciones para que la persona elija. El modelo "
     "propone; nunca escribe solo."),
    ("¿Esto reemplaza a My Inventory?", "DIANA",
     "No. Lo alimenta. Hoy comparamos contra el extracto que ustedes nos dieron; "
     "el siguiente paso es la integración directa."),
    ("¿Cuánto cuesta?", "DIANA",
     "Infraestructura AWS y consumo de IA por uso. Sin licencias por usuario."),
    ("¿Por qué la exactitud es del 88%?", "LUIS",
     "Porque 12 de cada 100 referencias estaban descuadradas ANTES de que "
     "llegáramos. Ese número no mide qué tan bien contamos: mide qué tan lejos "
     "estaba el sistema de la bodega. Encontrarlo es el valor."),
    ("¿Es accesible?", "LUIS",
     "Sí, y está medido: cero incumplimientos WCAG A/AA en las 17 pantallas, con "
     "axe-core. Navegación completa por teclado, alto contraste y tamaño de letra "
     "ajustable."),
]

ANTES = [
    "Tableta con la sesión de LUIS ya abierta en Conteo, bodega SIN abrir.",
    "Segundo dispositivo o pestaña con DIANA en Auditoría.",
    "Probar el micrófono UNA VEZ con la frase de la demo.",
    "Volumen del computador al máximo: el agente responde hablando.",
    "A la mano: d13g9u0pgpag0b.cloudfront.net · luis / diana · StockXperts1",
]


def mmss(s):
    return "%d:%02d" % (s // 60, s % 60)


# ─────────────────────── Word ───────────────────────
def parrafo(doc, texto, tam=10.5, color=None, negrita=False, cursiva=False,
            antes=0, despues=4, izq=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    if izq:
        p.paragraph_format.left_indent = Cm(izq)
    r = p.add_run(texto)
    r.font.size = Pt(tam)
    r.font.bold = negrita
    r.font.italic = cursiva
    if color:
        r.font.color.rgb = color
    return p


def armar_docx():
    doc = Document()
    est = doc.styles["Normal"]
    est.font.name = "Calibri"
    est.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(1.8)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("Guion de presentación · CuentaVoz")
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = NAVY

    total = sum(g[3] for g in GUION)
    parrafo(doc, "%s exactos · %d diapositivas · Luis y Diana"
            % (mmss(total), len(GUION)), 12, AZUL, negrita=True, despues=10)

    parrafo(doc,
            "El reparto aprovecha algo que ya tienen a favor: la plataforma tiene dos "
            "perfiles y ustedes son dos. Luis presenta y demuestra como auxiliar de "
            "inventarios; Diana cierra el ciclo como administradora de bodega. No es un "
            "truco de exposición: es la operación real, contada por las dos personas que "
            "la viven.", 10, GRAFITO, cursiva=True, despues=14)

    parrafo(doc, "ANTES DE EMPEZAR (5 minutos antes)", 11, AZUL, negrita=True,
            antes=6, despues=6)
    for x in ANTES:
        parrafo(doc, "☐  " + x, 10.5, despues=3, izq=0.5)

    doc.add_page_break()

    acumulado = 0
    for numero, titulo, quien, seg, lineas, notas in GUION:
        acumulado += seg
        cab = doc.add_paragraph()
        cab.paragraph_format.space_before = Pt(12)
        cab.paragraph_format.space_after = Pt(2)
        a = cab.add_run("%d · %s" % (numero, titulo))
        a.font.size = Pt(14)
        a.font.bold = True
        a.font.color.rgb = NAVY
        b = cab.add_run("     %s     %s  (acumulado %s)"
                        % (quien, mmss(seg), mmss(acumulado)))
        b.font.size = Pt(10)
        b.font.bold = True
        b.font.color.rgb = AZUL

        for l in lineas:
            if not l:
                continue
            parrafo(doc, l, 11, despues=5, izq=0.4)
        for n in notas:
            parrafo(doc, "▸ " + n, 9.5, AMARILLO_TX, cursiva=True, despues=3, izq=0.4)

    doc.add_page_break()
    parrafo(doc, "Control de tiempo", 16, NAVY, negrita=True, despues=8)
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = "Light Grid Accent 1"
    tabla.rows[0].cells[0].text = "Corte"
    tabla.rows[0].cells[1].text = "Deberían ir en"
    ac = 0
    cortes = {5: "Terminando la 5", 6: "Saliendo de la demo",
              10: "Terminando la 10", 13: "Empezando la 14"}
    for numero, _t, _q, seg, _l, _n in GUION:
        ac += seg
        if numero in cortes:
            f = tabla.add_row().cells
            f[0].text = cortes[numero]
            f[1].text = mmss(ac)
    parrafo(doc,
            "Si a los 5:00 no han salido de la demo, córtenla y sigan. La demo es lo más "
            "fuerte, pero quedarse sin llegar a la 14 es perder el pedido.",
            10, GRAFITO, cursiva=True, antes=8)

    parrafo(doc, "Preguntas que probablemente hagan", 16, NAVY, negrita=True,
            antes=16, despues=8)
    for pregunta, quien, respuesta in PREGUNTAS:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run("«%s»" % pregunta)
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = NAVY
        r2 = p.add_run("     " + quien)
        r2.font.size = Pt(9.5)
        r2.font.bold = True
        r2.font.color.rgb = AZUL
        parrafo(doc, respuesta, 10.5, despues=4, izq=0.4)

    doc.save(DOCX)
    return total


# ─────────────────── notas del PowerPoint ───────────────────
def poner_notas():
    pres = Presentation(PPTX)
    if len(pres.slides) != len(GUION):
        print("  aviso: %d diapositivas y %d entradas de guion"
              % (len(pres.slides), len(GUION)))
    acumulado = 0
    puestas = 0
    for (numero, titulo, quien, seg, lineas, notas), hoja in zip(GUION, pres.slides):
        acumulado += seg
        partes = ["%s  ·  %s  ·  acumulado %s" % (quien, mmss(seg), mmss(acumulado)), ""]
        partes += [l for l in lineas]
        if notas:
            partes += [""] + ["> " + n for n in notas]
        hoja.notes_slide.notes_text_frame.text = "\n".join(partes)
        puestas += 1
    pres.save(PPTX)
    return puestas


total = armar_docx()
print("  Word:  %s  (%s, %d laminas)" % (os.path.basename(DOCX), mmss(total), len(GUION)))
print("  PPTX:  %d diapositivas con notas del orador" % poner_notas())
