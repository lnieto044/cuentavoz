# -*- coding: utf-8 -*-
"""Regenera el codigo QR de la ultima diapositiva.

El QR es una IMAGEN: cambiar la URL escrita al lado no lo actualiza. El de
la presentacion seguia apuntando a https://cuentavoz.onrender.com, o sea
que quien lo escaneara durante la presentacion aterrizaba en el despliegue
viejo de Render en vez de en la plataforma.

Es el tipo de cosa que no se ve revisando el texto -el QR se ve igual
apunte a donde apunte- y que solo aparece decodificandolo. Por eso este
guion decodifica el resultado antes de darlo por bueno.

    python docs/actualizar_qr.py
"""
import io
import os
import sys

import cv2
import numpy as np
import qrcode
from pptx import Presentation
from pptx.util import Emu

AQUI = os.path.dirname(os.path.abspath(__file__))
URL = "https://d13g9u0pgpag0b.cloudfront.net/"
LADO = 372          # el mismo tamano en pixeles que traia el QR anterior

ARCHIVOS = [
    os.path.join(AQUI, "CuentaVoz_Colsubsidio_V2.pptx"),
    os.path.join(AQUI, "CuentaVoz_Colsubsidio_V2_conGIF.pptx"),
]


def hacer_qr(url, lado):
    """Negro sobre blanco y con un borde minimo, como el que habia."""
    q = qrcode.QRCode(version=None, box_size=10, border=1,
                      error_correction=qrcode.constants.ERROR_CORRECT_M)
    q.add_data(url)
    q.make(fit=True)
    im = q.make_image(fill_color="black", back_color="white").convert("RGB")
    im = im.resize((lado, lado))
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def decodificar(datos):
    arr = cv2.imdecode(np.frombuffer(datos, np.uint8), cv2.IMREAD_COLOR)
    texto, _, _ = cv2.QRCodeDetector().detectAndDecode(arr)
    return texto


nuevo = hacer_qr(URL, LADO)
leido = decodificar(nuevo)
if leido != URL:
    print("  el QR generado no se lee como esperaba: %r" % leido)
    sys.exit(1)
print("  QR generado y verificado -> %s" % leido)

for ruta in ARCHIVOS:
    if not os.path.exists(ruta):
        continue
    pres = Presentation(ruta)
    hoja = pres.slides[-1]

    # el QR es la imagen cuadrada grande de la ultima lamina: los demas
    # dibujos son iconos de 0,28" y el logo, que es apaisado
    objetivo = None
    for f in hoja.shapes:
        if f.__class__.__name__ != "Picture":
            continue
        if f.width == f.height and f.width > Emu(int(1.0 * 914400)):
            objetivo = f
            break
    if objetivo is None:
        print("  %s: no encontre el QR" % os.path.basename(ruta))
        continue

    viejo = decodificar(objetivo.image.blob)
    izq, arr_, anc, alt = objetivo.left, objetivo.top, objetivo.width, objetivo.height
    objetivo._element.getparent().remove(objetivo._element)

    tmp = os.path.join(AQUI, "_qr_tmp.png")
    io.open(tmp, "wb").write(nuevo)
    hoja.shapes.add_picture(tmp, izq, arr_, anc, alt)
    os.remove(tmp)
    pres.save(ruta)
    print("  %-38s  %s  ->  %s"
          % (os.path.basename(ruta), viejo or "(ilegible)", URL))

# comprobacion final sobre lo guardado
print()
for ruta in ARCHIVOS:
    if not os.path.exists(ruta):
        continue
    pres = Presentation(ruta)
    for f in pres.slides[-1].shapes:
        if (f.__class__.__name__ == "Picture" and f.width == f.height
                and f.width > Emu(int(1.0 * 914400))):
            print("  %-38s lee: %s" % (os.path.basename(ruta),
                                       decodificar(f.image.blob)))
